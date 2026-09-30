"""Generates the official B.Tech 3rd Year SRS Document (.docx) conforming to IEEE Std 830-1998.

Populates all academic sections, styled tables, clickable live links, and embedded high-resolution diagrams.
Candidate: Narayan Singh | Roll No: 2400320100734 | ABES EC / AKTU Lucknow
"""

import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from twin.paths import PROJECT_ROOT, DOCS_DIR

DOCX_OUTPUT = PROJECT_ROOT / "BTech_3rd_Year_SRS_Final.docx"
FIG_DIR = DOCS_DIR / "phase3_demo" / "figures"


def set_cell_background(cell, fill_hex):
    """Set cell shading color."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_hyperlink(paragraph, url, text, color="1E40AF", underline=True):
    """Places a clickable hyperlink into a python-docx paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._element.append(hyperlink)


def style_table(table, col_widths=None):
    """Applies clean engineering table borders, headers, and zebra-striping."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        if i == 0:
            for cell in row.cells:
                set_cell_background(cell, "1E3A8A")  # Deep Royal Navy
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.5)
        else:
            bg_color = "F8FAFC" if i % 2 == 1 else "FFFFFF"
            for cell in row.cells:
                set_cell_background(cell, bg_color)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.0)
                        r.font.color.rgb = RGBColor(30, 41, 59)

        for j, cell in enumerate(row.cells):
            set_cell_margins(cell, top=110, bottom=110, left=130, right=130)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]


def add_figure_with_caption(doc, fig_filename, caption_text, width_in=6.2):
    """Embeds high-resolution figure and academic caption cleanly."""
    fig_path = FIG_DIR / fig_filename
    if fig_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(3)
        doc.add_picture(str(fig_path), width=Inches(width_in))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.italic = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = RGBColor(71, 85, 105)


def build_srs_document():
    doc = Document()

    # Standard 1.0 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Typography
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # --------------------------------------------------------------------------
    # COVER PAGE
    # --------------------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("DEPARTMENT OF MECHANICAL ENGINEERING\nABES ENGINEERING COLLEGE / AKTU, LUCKNOW\n")
    r_inst.font.bold = True
    r_inst.font.size = Pt(13)
    r_inst.font.color.rgb = RGBColor(30, 58, 138)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(20)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_type = p_title.add_run("SOFTWARE REQUIREMENTS SPECIFICATION (SRS)\n")
    r_type.font.bold = True
    r_type.font.size = Pt(18)
    r_type.font.color.rgb = RGBColor(15, 23, 42)

    r_proj = p_title.add_run("Knowledge-Driven IoT Fault Diagnosis Assistant:\nAI-Driven Digital Twin for Smart Factory Operations\n")
    r_proj.font.bold = True
    r_proj.font.size = Pt(16)
    r_proj.font.color.rgb = RGBColor(37, 99, 235)

    r_std = p_title.add_run("Standard Academic Compliance: IEEE Std 830-1998\n")
    r_std.font.italic = True
    r_std.font.size = Pt(11)
    r_std.font.color.rgb = RGBColor(100, 116, 139)

    p_divider = doc.add_paragraph()
    p_divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_divider.paragraph_format.space_after = Pt(25)
    r_div = p_divider.add_run("—" * 45)
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    # Submission Box
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.line_spacing = 1.3
    r_sub = p_meta.add_run("SUBMITTED BY:\n")
    r_sub.font.bold = True
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(30, 58, 138)

    p_meta.add_run("Candidate Name: ")
    r_name = p_meta.add_run("Narayan Singh\n")
    r_name.font.bold = True
    p_meta.add_run("University Roll No: ")
    r_roll = p_meta.add_run("2400320100734\n")
    r_roll.font.bold = True
    p_meta.add_run("Degree: Bachelor of Technology (B.Tech) in Mechanical Engineering\n")
    p_meta.add_run("Semester / Year: 6th Semester (3rd Year, 2026 – 2027)\n\n")

    r_gui = p_meta.add_run("UNDER THE GUIDANCE OF:\n")
    r_gui.font.bold = True
    r_gui.font.size = Pt(11)
    r_gui.font.color.rgb = RGBColor(30, 58, 138)
    p_meta.add_run("Supervisor: ")
    r_guide = p_meta.add_run("Prof. / Dr. Tanvi Saxena\n")
    r_guide.font.bold = True
    p_meta.add_run("Department: Department of Mechanical Engineering\n")
    p_meta.add_run("Institution: ABES Engineering College, Ghaziabad / AKTU, Lucknow\n")
    p_meta.add_run("Date of Submission: 30-September-2026\n\n")

    r_lnk = p_meta.add_run("LIVE PROJECT REPOSITORIES & REMOTE PORTS:\n")
    r_lnk.font.bold = True
    r_lnk.font.size = Pt(10.5)
    r_lnk.font.color.rgb = RGBColor(30, 58, 138)

    p_link1 = doc.add_paragraph()
    p_link1.add_run("• Public Grafana Dashboard: ")
    add_hyperlink(p_link1, "https://homeless-interfaces-experienced-cio.trycloudflare.com", "https://homeless-interfaces-experienced-cio.trycloudflare.com (Cloudflare Tunnel)")

    p_link2 = doc.add_paragraph()
    p_link2.add_run("• GitHub Source Repository: ")
    add_hyperlink(p_link2, "https://github.com/Narayan1006/Factory_Automation.git", "https://github.com/Narayan1006/Factory_Automation.git")

    p_link3 = doc.add_paragraph()
    p_link3.add_run("• Interactive Web Console: Streamlit Cloud / Local Engine (`app.py`)\n")

    doc.add_page_break()

    # --------------------------------------------------------------------------
    # REVISION HISTORY TABLE
    # --------------------------------------------------------------------------
    h_rev = doc.add_heading("Document Revision & Approval History", level=2)
    h_rev.style.font.color.rgb = RGBColor(30, 58, 138)

    table_rev = doc.add_table(rows=4, cols=5)
    headers_rev = ["Version", "Date", "Description / Major Changes", "Prepared By", "Reviewed By"]
    for j, h in enumerate(headers_rev):
        table_rev.rows[0].cells[j].paragraphs[0].text = h

    rev_data = [
        ("0.1", "15-Sep-2026", "Initial Draft, Problem Definition & Scope Boundary", "Narayan Singh", "Dr. Tanvi Saxena"),
        ("0.5", "22-Sep-2026", "Multi-tier Architecture, InfluxDB Line Protocol & PyTorch MLP", "Narayan Singh", "Dr. Tanvi Saxena"),
        ("1.0", "30-Sep-2026", "Final 3rd-Year SRS Baseline, Cloudflare Tunnel & IEEE 830 Compliance", "Narayan Singh", "HOD / Review Committee"),
    ]
    for i, row in enumerate(rev_data):
        for j, val in enumerate(row):
            table_rev.rows[i + 1].cells[j].paragraphs[0].text = val
    style_table(table_rev, [Inches(0.8), Inches(1.1), Inches(2.7), Inches(1.3), Inches(1.3)])

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # --------------------------------------------------------------------------
    # 1. INTRODUCTION
    # --------------------------------------------------------------------------
    h1 = doc.add_heading("1. Introduction", level=1)
    h1.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_heading("1.1 Purpose", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "This Software Requirements Specification (SRS) establishes the complete functional and non-functional requirements "
        "for the project 'Knowledge-Driven IoT Fault Diagnosis Assistant: AI-Driven Digital Twin for Smart Factory Operations'. "
        "Developed under the academic curriculum of the Department of Mechanical Engineering, ABES Engineering College / AKTU Lucknow, "
        "this specification strictly adheres to IEEE Std 830-1998 standards for software engineering documentation. The intended "
        "audience comprises academic evaluators, university project committees, industrial mentors, and smart manufacturing engineers."
    )

    doc.add_heading("1.2 Scope of the System", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "The software product is a cyber-physical industrial Digital Twin modeled on Line 3 of the Bosch Production Line Performance "
        "dataset (1.18 million parts, 52 stations). The digital twin tracks parts moving along the core sequence S29 -> S30 -> S33 -> S34 -> "
        "(S35 or S36) -> S37, which represents 89.02% of all factory production. The primary objective is to eliminate wasteful late-stage "
        "defects by providing 30 to 75 minutes of proactive quality lead time before parts reach physical exit inspection at S37."
    )

    add_figure_with_caption(
        doc,
        "figure3_lead_time_comparison.png",
        "Figure 1.1: Early Quality Warning Lead Time Comparison (+30 to 75 min Saved)",
        width_in=6.0,
    )

    p_scope = doc.add_paragraph()
    r = p_scope.add_run("In-Scope Capabilities:\n")
    r.font.bold = True
    p_scope.add_run(
        "• Deterministic shop-floor telemetry replay over Eclipse Mosquitto MQTT with speed scaling (1x to 120x, burst >7,000 events/s).\n"
        "• Real-time topological tracking of Line 3 core flow handling 89.02% of total factory volume (1,053,742 parts).\n"
        "• Stateful multi-factor Station Health Scoring (Hs in [0, 100]) combining sensor Z-score drift, pacing delay, and defect rates.\n"
        "• PyTorch Multilayer Perceptron (MLP) defect risk prediction evaluated at the S34 fork with Platt probability calibration.\n"
        "• 1-layer PyTorch LSTM Autoencoder modeling 24-hour multivariate window sequences for unsupervised process anomaly detection.\n"
        "• Sub-second time-series persistence in InfluxDB v2 and live operational visualization via Grafana 11 and Streamlit.\n"
        "• Cloudflare Zero-Trust Secure Tunneling for global public dashboard access with anonymous viewer mode."
    )

    doc.add_heading("1.3 Definitions, Acronyms, and Abbreviations", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    table_def = doc.add_table(rows=8, cols=3)
    def_headers = ["Category", "Term / Acronym", "Definition / Standard Industrial Context"]
    for j, h in enumerate(def_headers):
        table_def.rows[0].cells[j].paragraphs[0].text = h
    defs = [
        ("Standard", "SRS", "Software Requirements Specification conforming to IEEE Std 830-1998 format."),
        ("Architecture", "Digital Twin", "Virtual software representation of physical machines and factory processes updated in real time."),
        ("Network", "MQTT", "Message Queuing Telemetry Transport (ISO/IEC 20922); industrial lightweight pub/sub protocol."),
        ("Database", "TSM / Flux", "Time-Structured Merge tree storage engine in InfluxDB v2; Flux functional data query language."),
        ("Deep Learning", "MLP / LSTM-AE", "Multilayer Perceptron for defect classification; LSTM Autoencoder for temporal anomaly detection."),
        ("Statistics", "PR-AUC", "Precision-Recall Area Under Curve; robust performance metric under severe class imbalance (0.51% defect rate)."),
        ("Security", "Cloudflare Tunnel", "Outbound zero-trust reverse proxy exposing local operational services securely to the internet."),
    ]
    for i, row in enumerate(defs):
        for j, val in enumerate(row):
            table_def.rows[i + 1].cells[j].paragraphs[0].text = val
    style_table(table_def, [Inches(1.2), Inches(1.3), Inches(4.5)])

    # --------------------------------------------------------------------------
    # 2. OVERALL DESCRIPTION
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h2 = doc.add_heading("2. Overall Description", level=1)
    h2.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_heading("2.1 Product Perspective & Cyber-Physical Architecture", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "The digital twin functions as a multi-tier industrial edge intelligence platform. The physical layer (Bosch Line 3) "
        "is deterministically emulated by the Historical Replay Engine over MQTT. The data layer ingests sub-second telemetry "
        "into InfluxDB v2 TSM storage. The intelligence layer executes PyTorch defect predictions and LSTM autoencoder reconstructions "
        "in real time. The visualization layer presents dual operational interfaces: Grafana 11 for industrial SCADA operations "
        "and Streamlit for interactive scenario exploration."
    )

    add_figure_with_caption(
        doc,
        "figure5_system_architecture.png",
        "Figure 2.1: 4-Tier End-to-End Cyber-Physical System Architecture",
        width_in=6.2,
    )

    add_figure_with_caption(
        doc,
        "figure1_line3_topology.png",
        "Figure 2.2: Bosch Line 3 Core Manufacturing Topology & Informative Feature Allocation",
        width_in=6.2,
    )

    doc.add_heading("2.2 User Classes and Characteristics", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    table_users = doc.add_table(rows=5, cols=3)
    u_headers = ["User Role / Persona", "Technical Proficiency", "Core Responsibilities & Access Rights"]
    for j, h in enumerate(u_headers):
        table_users.rows[0].cells[j].paragraphs[0].text = h
    users_data = [
        ("Plant Operations Engineer", "High (Industrial Engg)", "Monitors Line Overview; tracks plant WIP, station health matrix, and line throughput."),
        ("Quality Control Specialist", "Intermediate (QA / QC)", "Monitors Parts & Risk and AI Insights; inspects parts with risk > 0.70; reviews early warnings."),
        ("Maintenance Technician", "Intermediate (Shop Floor)", "Responds to station health warnings (Hs < 50); diagnoses sensor feature drift and transit delays."),
        ("Academic Evaluator / Guide", "High (Software / AI Research)", "Audits system architecture, causal validation integrity, ML metrics, and test coverage."),
    ]
    for i, row in enumerate(users_data):
        for j, val in enumerate(row):
            table_users.rows[i + 1].cells[j].paragraphs[0].text = val
    style_table(table_users, [Inches(1.8), Inches(1.5), Inches(3.7)])

    doc.add_heading("2.3 Operating Environment & Technical Stack", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "• Host Hardware: Intel Core i5/i7/i9 (x86_64) CPU, 16 GB RAM minimum.\n"
        "• Hardware Acceleration: NVIDIA GeForce RTX 3050 Laptop GPU (6 GB VRAM) with CUDA 12.4 support (automatic CPU fallback supported).\n"
        "• Host Operating System: Microsoft Windows 11 / Windows 10 (PowerShell 7+) or Ubuntu 22.04 LTS.\n"
        "• Virtualization & Networking: Docker Desktop 4.28+ with Linux container subsystem and Cloudflare Tunnel.\n"
        "• Core Software Frameworks: Python 3.12, PyTorch 2.6.0, InfluxDB v2.7, Eclipse Mosquitto 2.0, Grafana 11.1.0, Streamlit 1.35+."
    )

    doc.add_heading("2.4 Design and Implementation Constraints", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "1. Zero Data Fabrication: Machine identifiers and feature names strictly preserve empirical Bosch nomenclature (S29-S37).\n"
        "2. Strict Temporal Causality: Part defect risk is evaluated strictly at the entry timestamp of S35/S36. No future timestamps from S37 are leaked.\n"
        "3. Imbalanced Learning: The defect prevalence is strictly 0.51% (1 defect in 196 parts). Accuracy is strictly prohibited as a metric; PR-AUC, MCC, and Lift@1% are mandatory.\n"
        "4. Deterministic Reproducibility: Replay engine uses fixed random seeds and pre-computed topological routing matrices."
    )

    add_figure_with_caption(
        doc,
        "figure2_scenario_timelines.png",
        "Figure 2.3: Leave-Scenario-Out Empirical Validation Timeline across 1,700 Sim Hours",
        width_in=6.2,
    )

    # --------------------------------------------------------------------------
    # 3. SYSTEM FEATURES AND REQUIREMENTS
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h3 = doc.add_heading("3. System Features & Functional Requirements", level=1)
    h3.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_heading("3.1 Module 1: Ingestion, Streaming & Discrete-Event Replay", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "The Replay Engine orchestrates a virtual discrete-event clock advancing at user-selected speeds (1x to 120x, or burst mode >7,000 events/s). "
        "It sequentially reads Parquet tables and publishes MQTT JSON payloads to topics structured as 'factory/line3/{station_id}/telemetry'. "
        "The Ingestion Daemon subscribes to all line topics, parses telemetry, and writes batches to InfluxDB v2 using line protocol."
    )

    doc.add_heading("3.2 Module 2: Edge Analytics, Health Scoring & AI Inference", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "The Scoring Service processes station sensor telemetry through two complementary paradigms:\n"
        "1. Multi-Factor Station Health Scoring: A deterministic composite index Hs = 100 - (0.40 * S_drift + 0.35 * S_pacing + 0.25 * S_defect).\n"
        "2. Calibrated PyTorch Deep Learning: A 3-layer MLP (128-64 architecture, ReLU, Dropout 0.3, weighted BCE) with Platt scaling "
        "that predicts quality risk in 11.4 milliseconds, and a PyTorch LSTM Autoencoder reconstructing 24-hour multivariate window sequences."
    )

    add_figure_with_caption(
        doc,
        "figure8_health_score_breakdown.png",
        "Figure 3.1: Multi-Factor Station Health Score Breakdown (Nominal vs Anomaly State)",
        width_in=6.2,
    )

    add_figure_with_caption(
        doc,
        "figure6_deep_learning_models.png",
        "Figure 3.2: Dual Deep Learning Architectures (Defect MLP + LSTM Autoencoder)",
        width_in=6.2,
    )

    add_figure_with_caption(
        doc,
        "figure7_pr_curve_and_confusion.png",
        "Figure 3.3: Precision-Recall Evaluation Curves & Confusion Matrix (Lift = 7.26x)",
        width_in=6.2,
    )

    add_figure_with_caption(
        doc,
        "figure4_risk_bands_overlay.png",
        "Figure 3.4: Station S34 Cumulative Risk Bands vs Empirical Defect Prevalence",
        width_in=6.2,
    )

    add_figure_with_caption(
        doc,
        "figure9_feature_importance.png",
        "Figure 3.5: Global Permutation Feature Importance Across 85 Predictive Features",
        width_in=6.2,
    )

    # --------------------------------------------------------------------------
    # 4. EXTERNAL INTERFACES & CLOUD DEPLOYMENT
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h4 = doc.add_heading("4. External Interfaces & Cloud Architecture", level=1)
    h4.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_heading("4.1 User Interfaces (Industrial SCADA & Cloud Web Console)", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "• Grafana 11 Industrial SCADA: 4 provisioned dashboards (Line Overview, Station Detail, Parts & Risk Tracking, and AI Insights) "
        "querying InfluxDB v2 via Flux. Features eliminated tag explosion and clean scalar matrix status.\n"
        "• Streamlit Cloud Web Console: An interactive minimalist control center (black, charcoal, brushed silver) "
        "allowing external reviewers to trigger 1-click scenario simulations, stream real-time events, and test PyTorch inference live."
    )

    doc.add_heading("4.2 Cloud Deployment & Zero-Trust Remote Accessibility", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph(
        "To allow external academic evaluators to inspect the running digital twin without requiring local installations or port forwarding, "
        "the project incorporates a two-channel cloud accessibility architecture:\n\n"
        "1. Streamlit Community Cloud: Deployed via GitHub repository synchronization (app.py and root requirements.txt), serving "
        "the interactive simulation lab globally over HTTPS.\n"
        "2. Cloudflare Zero-Trust Secure Tunnel (cloudflared): Exposes local Grafana port 3000 to the global internet via an encrypted "
        "outbound tunnel on trycloudflare.com. Anonymous Viewer mode (GF_AUTH_ANONYMOUS_ENABLED=true) enables instant 1-click evaluation "
        "without credentials.\n\n"
        "Active Public Verification Endpoints:"
    )

    p_live_ep = doc.add_paragraph()
    p_live_ep.add_run("• Public Grafana Dashboard: ")
    add_hyperlink(p_live_ep, "https://homeless-interfaces-experienced-cio.trycloudflare.com", "https://homeless-interfaces-experienced-cio.trycloudflare.com")
    p_live_ep.add_run("\n• Public GitHub Repository: ")
    add_hyperlink(p_live_ep, "https://github.com/Narayan1006/Factory_Automation.git", "https://github.com/Narayan1006/Factory_Automation.git")
    p_live_ep.add_run("\n• InfluxDB Data Engine: http://localhost:8086 (Bucket: factory_telemetry, Org: bosch_twin)\n")
    p_live_ep.add_run("• MQTT Broker: localhost:1883 (Topics: factory/line3/#)")

    # --------------------------------------------------------------------------
    # 5. NON-FUNCTIONAL REQUIREMENTS (NFRs)
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h5 = doc.add_heading("5. Non-Functional Requirements (NFRs)", level=1)
    h5.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_heading("5.1 Quantitative Benchmarks & Validation Matrix", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    table_nfr = doc.add_table(rows=7, cols=4)
    nfr_headers = ["NFR ID", "Quality Attribute", "Target Metric / Standard", "Empirical Validation Result"]
    for j, h in enumerate(nfr_headers):
        table_nfr.rows[0].cells[j].paragraphs[0].text = h

    nfr_data = [
        ("NFR-1", "Performance (Throughput)", "Replay engine sustains >= 5,000 events/sec in burst mode.", "Achieved 7,405 events/sec (0 drops)."),
        ("NFR-2", "Performance (Latency)", "Real-time GPU neural inference < 25 ms per part.", "Achieved 11.4 ms on RTX 3050 GPU."),
        ("NFR-3", "Security & Access", "Token-based DB auth; zero plaintext passwords; RBAC.", "Verified in Docker & twin_config.yaml."),
        ("NFR-4", "Reliability & Uptime", "Zero memory leakage across 24h streaming simulation.", "Deque-bounded windows (maxlen=10,000)."),
        ("NFR-5", "ML Precision Lift", "Defect model Lift@1% >= 5.0x over random baseline.", "Achieved 5.8x to 7.26x Lift@1%."),
        ("NFR-6", "Portability", "1-Click multi-container launch; reproducible wheel.", "Verified via scripts/demo.ps1 runbook."),
    ]
    for i, row in enumerate(nfr_data):
        for j, val in enumerate(row):
            table_nfr.rows[i + 1].cells[j].paragraphs[0].text = val
    style_table(table_nfr, [Inches(0.9), Inches(1.8), Inches(2.2), Inches(2.1)])

    # --------------------------------------------------------------------------
    # 6. VERIFICATION, USE CASES & DFD
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h6 = doc.add_heading("6. System Design, Use Cases & Data Flow Architecture", level=1)
    h6.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_heading("6.1 Use Case Specifications", level=2).style.font.color.rgb = RGBColor(30, 58, 138)
    table_uc = doc.add_table(rows=6, cols=4)
    uc_headers = ["Use Case ID", "Use Case Name", "Primary Actor", "Operational Description"]
    for j, h in enumerate(uc_headers):
        table_uc.rows[0].cells[j].paragraphs[0].text = h
    uc_data = [
        ("UC-01", "Launch Replay & Set Speed", "Plant Operations Eng.", "Initiates scenario streaming at configurable speed multiplier."),
        ("UC-02", "Monitor Line Overview", "Plant Operations Eng.", "Observes throughput, WIP, and real-time Station Status Matrix."),
        ("UC-03", "Inspect Station Health & Drift", "Maintenance Tech.", "Analyzes feature Z-score deviations and transit pacing percentiles."),
        ("UC-04", "Predictive Defect Screening", "Quality Specialist", "Evaluates parts at S34 fork; flags high-risk units (+30-75m lead time)."),
        ("UC-05", "LSTM Anomaly Investigation", "Quality Specialist", "Drills into reconstruction error spikes and station contribution radar."),
    ]
    for i, row in enumerate(uc_data):
        for j, val in enumerate(row):
            table_uc.rows[i + 1].cells[j].paragraphs[0].text = val
    style_table(table_uc, [Inches(1.0), Inches(1.8), Inches(1.6), Inches(2.6)])

    add_figure_with_caption(
        doc,
        "figure10_use_case_and_dfd.png",
        "Figure 6.1: High-Level Use Case Map & DFD Level 0 Context Architecture",
        width_in=6.2,
    )

    # --------------------------------------------------------------------------
    # 7. ACADEMIC REVIEW & SIGN-OFF
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h7 = doc.add_heading("7. Academic Review & Sign-Off", level=1)
    h7.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "This Software Requirements Specification document has been submitted and reviewed by the academic supervisory "
        "panel at ABES Engineering College / AKTU Lucknow as the binding technical specification for the B.Tech 3rd-year engineering project:\n\n"
        "Project Title: Knowledge-Driven IoT Fault Diagnosis Assistant: AI-Driven Digital Twin for Smart Factory Operations\n"
        "Candidate Name: Narayan Singh (University Roll No: 2400320100734)\n"
        "Degree: Bachelor of Technology (B.Tech) in Mechanical Engineering\n"
        "Academic Year: 2026 – 2027\n"
    )

    doc.add_paragraph().paragraph_format.space_before = Pt(20)

    # Signature Table
    sig_table = doc.add_table(rows=4, cols=3)
    sig_headers = ["Project Supervisor / Guide", "Internal Examiner", "Head of Department (HOD)"]
    for j, h in enumerate(sig_headers):
        sig_table.rows[0].cells[j].paragraphs[0].text = h

    sig_data = [
        ("Prof. / Dr. Tanvi Saxena\nDept. of Mechanical Engineering", "Evaluation Committee\nDept. of Mechanical Engineering", "Head of Department\nDept. of Mechanical Engineering"),
        ("\nSignature: __________________\n", "\nSignature: __________________\n", "\nSignature: __________________\n"),
        ("Date: ________________________", "Date: ________________________", "Date: ________________________"),
    ]
    for i, row in enumerate(sig_data):
        for j, val in enumerate(row):
            sig_table.rows[i + 1].cells[j].paragraphs[0].text = val

    style_table(sig_table, [Inches(2.3), Inches(2.3), Inches(2.3)])

    doc.save(str(DOCX_OUTPUT))
    print(f"Document successfully created: {DOCX_OUTPUT}")


if __name__ == "__main__":
    build_srs_document()
